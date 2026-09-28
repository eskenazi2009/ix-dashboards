# -*- coding: utf-8 -*-
"""
Ventas diarias de las 3 tiendas IX (RB Guatemala, NB El Salvador, NB Dominicana)
SIN depender de la PC. Lo usa la rutina de Claude en la nube (9 pm Panama) y
tambien sirve a mano.

  python cloud/sync_ix.py pedir  --desde YYYY-MM-DD --hasta YYYY-MM-DD [--cuentas GT,SV,RD]
      Pide al portal TotalCommerce el export de ventas de cada cuenta. El portal
      contesta por correo ("Descargable de Ventas", totalcommerce@ixcomercio.com)
      con un enlace .xlsx que dura 5 dias. El correo NO dice el pais: se sabe por
      el contenido del Excel.

  python cloud/sync_ix.py cargar URL_O_ARCHIVO [URL_O_ARCHIVO ...]
      Baja cada export, detecta su pais y su rango de fechas, suma venta neta y
      unidades por dia y lo mete en ventas-ix.enc (historial cifrado); luego
      reemplaza las tiendas IX dentro de ventas-macro.enc (datos de la app).
      Los dias que cubre el export se re-escriben (gana el export mas nuevo).
      Requiere IX_CLAVE en el entorno. Nada queda en claro en el repo.

Columnas del export de 44 columnas (0-based): 2 canal, 4 pais, 7 tienda, 8 fecha,
28 unidades, 39 neto. El neto es sin impuestos; RD viene ya en USD.
"""
import datetime, hashlib, io, json, os, re, secrets, sys
import requests
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.dirname(HERE)

API = "https://ca-send-sales-portal-prod.ambitiousplant-fd82aa00.eastus2.azurecontainerapps.io/api/sales/raymond@bmc.com.pa?filetype=xlsx"
HDR = {"Accept": "application/json, text/plain, */*", "x-api-version": "1", "x-channel": "DIG,PH",
       "x-customerid": "1", "Origin": "https://www.totalcommerce.pro", "Referer": "https://www.totalcommerce.pro/"}
CUENTAS = {  # id -> (x-country, x-commerce, texto en la columna Pais, brand, country para la app)
    "GT": ("GT", "NEB,RUN", "guatemala", "Running Balboa", "Guatemala"),
    "SV": ("SV", "NEB", "salvador", "New Balance", "El Salvador"),
    "RD": ("DO", "NEB", "dominic", "New Balance", "Dominicana"),
}
M = dict(canal=2, pais=4, tienda=7, fecha=8, units=28, net=39)

# ------------------------------------------------------------------ cifrado
def _aes():
    cfg = json.load(open(os.path.join(SITE, "crypt.json"), encoding="utf-8"))
    clave = os.environ.get("IX_CLAVE") or sys.exit("falta IX_CLAVE en el entorno")
    return AESGCM(hashlib.pbkdf2_hmac("sha256", clave.encode(), bytes.fromhex(cfg["salt"]), cfg["iter"], 32))
def dec(aes, path):
    b = open(path, "rb").read(); return json.loads(aes.decrypt(b[:12], b[12:], None).decode("utf-8"))
def enc(aes, path, obj):
    iv = secrets.token_bytes(12)
    open(path, "wb").write(iv + aes.encrypt(iv, json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode("utf-8"), None))

# ------------------------------------------------------------------ pedir
def pedir(desde, hasta, cuentas):
    for cid in cuentas:
        pais, com, *_ = CUENTAS[cid]
        h = dict(HDR, **{"x-country": pais, "x-commerce": com, "x-start": desde, "x-end": hasta})
        r = requests.get(API, headers=h, timeout=60)
        try: msg = r.json().get("message", r.text[:200])
        except Exception: msg = r.text[:200]
        print("pedido %s %s a %s -> HTTP %s: %s" % (cid, desde, hasta, r.status_code, msg))

# ------------------------------------------------------------------ cargar
def leer_export(src):
    """-> (cuenta, {fecha: [neto, uds]}, (dmin, dmax)) o None si no es un export valido."""
    if re.match(r"https?://", src):
        r = requests.get(src, timeout=120); r.raise_for_status(); data = r.content
    else:
        data = open(src, "rb").read()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rr = list(wb[wb.sheetnames[0]].iter_rows(values_only=True)); wb.close()
    if not rr or len([c for c in rr[0] if c is not None]) < 40:
        print("  OMITIDO (no es el export de 44 columnas):", src[-60:]); return None
    dias, paises = {}, {}
    for r in rr[1:]:
        if len(r) <= M["net"]: continue
        f = r[M["fecha"]]
        if not isinstance(f, datetime.datetime): continue
        txt = (str(r[M["pais"]] or "") + " " + str(r[M["tienda"]] or "")).lower()
        for cid, (_p, _c, key, *_r) in CUENTAS.items():
            if key in txt: paises[cid] = paises.get(cid, 0) + 1; break
        d = f.date().isoformat(); v = dias.setdefault(d, [0.0, 0])
        v[0] += float(r[M["net"]] or 0); v[1] += int(float(r[M["units"]] or 0))
    if not dias:
        print("  export sin ventas:", src[-60:]); return None
    cid = max(paises, key=paises.get) if paises else None
    if not cid: print("  OMITIDO: no pude identificar el pais:", src[-60:]); return None
    dias = {d: [round(v[0], 2), v[1]] for d, v in dias.items()}
    return cid, dias, (min(dias), max(dias))

def cargar(srcs, rango=None):
    aes = _aes()
    ix_path = os.path.join(SITE, "ventas-ix.enc")
    ix = dec(aes, ix_path) if os.path.exists(ix_path) else {"days": {}, "cov": {}}
    for src in srcs:
        res = leer_export(src)
        if not res: continue
        cid, dias, (a, b) = res
        # el rango pedido (si se paso) es la cobertura real: un dia sin filas = tienda cerrada
        ra, rb = rango if rango else (a, b)
        d = datetime.date.fromisoformat(ra)
        while d.isoformat() <= rb:
            k = d.isoformat()
            ix["days"].setdefault(k, {})[cid] = dias.get(k, [0.0, 0])
            d += datetime.timedelta(days=1)
        c = ix["cov"].get(cid)
        ix["cov"][cid] = [min(c[0], ra), max(c[1], rb)] if c else [ra, rb]
        print("  %s: %d dias con venta, cobertura %s a %s, neto $%s" % (cid, len(dias), ra, rb, format(sum(v[0] for v in dias.values()), ",.2f")))
    ix["days"] = dict(sorted(ix["days"].items()))
    enc(aes, ix_path, ix)
    # ---- app: reemplazar solo las tiendas IX
    macro_path = os.path.join(SITE, "ventas-macro.enc")
    macro = dec(aes, macro_path)
    macro["stores"] = [s for s in macro["stores"] if s.get("id") not in CUENTAS]
    for cid, (_p, _c, _k, brand, country) in CUENTAS.items():
        days = {d: v[cid] for d, v in ix["days"].items() if cid in v}
        if not days: continue
        cov = ix["cov"].get(cid) or [min(days), max(days)]
        macro["stores"].insert(0, dict(id=cid, brand=brand, country=country, cov=cov, days=days))
    macro["stores"].sort(key=lambda s: ["GT", "SV", "RD"].index(s["id"]) if s["id"] in CUENTAS else 9)
    macro["gen"] = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-5))).strftime("%Y-%m-%d %H:%M")
    enc(aes, macro_path, macro)
    for cid in CUENTAS:
        if cid in ix["cov"]: print("  cobertura %s: %s a %s" % (cid, *ix["cov"][cid]))

if __name__ == "__main__":
    a = sys.argv[1:]
    if a[:1] == ["pedir"]:
        desde = a[a.index("--desde") + 1]; hasta = a[a.index("--hasta") + 1]
        cuentas = a[a.index("--cuentas") + 1].split(",") if "--cuentas" in a else list(CUENTAS)
        pedir(desde, hasta, cuentas)
    elif a[:1] == ["cargar"]:
        rango = None
        if "--rango" in a:
            i = a.index("--rango"); rango = (a[i + 1], a[i + 2]); a = a[:i] + a[i + 3:]
        cargar(a[1:], rango)
    else:
        sys.exit(__doc__)
